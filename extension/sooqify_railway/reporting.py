"""CSV/XLSX report generation for daily Sooqify audits."""

from __future__ import annotations

import csv
import os
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from audit import audit_snapshot


TIMEZONE_NAME = os.getenv("APP_TIMEZONE", "Asia/Aden")
REPORT_TZ = ZoneInfo(TIMEZONE_NAME)
DETAIL_FIELDS = [
    "local_id", "store_product_id", "name_en", "name_ar", "store_name",
    "expected_brand", "expected_brand_id", "store_brand", "store_brand_id",
    "archive_price_sar", "expected_price_sar", "price_basis", "store_price_sar", "store_image_count", "match_basis", "store_fields_unavailable", "added_by",
    "workflow_status", "status", "findings",
]
EMPLOYEE_FIELDS = [
    "employee", "total", "issues", "unverified", "unresolved", "missing",
    "name_mismatch", "price_mismatch", "brand_mismatch", "store_id_mismatch",
    "no_image", "image_unknown",
]


def _parse_datetime(value):
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed
    except (TypeError, ValueError):
        return None


def _employee_rows(details):
    employees = defaultdict(lambda: {
        "employee": "", "total": 0, "issues": 0, "unverified": 0, "unresolved": 0,
        "missing": 0, "name_mismatch": 0, "price_mismatch": 0,
        "brand_mismatch": 0, "store_id_mismatch": 0, "no_image": 0, "image_unknown": 0,
    })
    for row in details:
        if row.get("local_id") is None:
            continue
        name = row.get("added_by") or "غير محدد"
        item = employees[name]
        item["employee"] = name
        item["total"] += 1
        findings = row.get("findings", [])
        item["issues"] += row.get("status") in {"issues", "missing"}
        item["unverified"] += row.get("status") == "unverified"
        item["unresolved"] += row.get("status") == "unlinked"
        item["missing"] += "missing_from_store" in findings
        item["name_mismatch"] += "name_mismatch" in findings
        item["no_image"] += "no_store_image" in findings
        item["image_unknown"] += "image_count_unavailable" in findings
        item["brand_mismatch"] += any(value in findings for value in ("brand_id_mismatch", "brand_name_mismatch"))
        item["store_id_mismatch"] += "store_id_mismatch" in findings
        item["price_mismatch"] += "store_price_mismatch" in findings
    return sorted(employees.values(), key=lambda row: (-row["issues"], row["employee"].casefold()))


def build_daily_report(snapshot, archive_items, price_baselines=None, now=None):
    if not isinstance(snapshot, dict) or snapshot.get("complete") is not True:
        raise ValueError("A complete Sooqify snapshot is required to create an audit report.")
    audit = audit_snapshot(snapshot, archive_items, price_baselines=price_baselines)
    now = now or datetime.now(REPORT_TZ)
    captured = _parse_datetime(snapshot.get("captured_at"))
    age_hours = round(max(0.0, (now.astimezone(timezone.utc) - captured).total_seconds() / 3600), 2) if captured else None
    result = {
        "generated_at": now.isoformat(timespec="seconds"),
        "timezone": TIMEZONE_NAME,
        "scan_id": snapshot.get("scan_id", ""),
        "scan_captured_at": snapshot.get("captured_at", ""),
        "scan_age_hours": age_hours,
        "scan_stale": age_hours is None or age_hours > 24,
        "scan_complete": True,
        "pages_scanned": snapshot.get("pages_scanned"),
        "snapshot_product_count": snapshot.get("product_count", len(snapshot.get("products", []))),
        "summary": audit["summary"],
        "details": audit["details"],
        "employees": _employee_rows(audit["details"]),
    }
    return result


def _safe_spreadsheet_value(value):
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="172554")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in sheet.columns:
        width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_report_files(report, reports_dir):
    reports_dir = Path(reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(REPORT_TZ).strftime("%Y%m%d_%H%M%S")
    csv_path = reports_dir / f"sooqify_audit_{timestamp}.csv"
    employee_csv_path = reports_dir / f"sooqify_audit_{timestamp}_employees.csv"
    xlsx_path = reports_dir / f"sooqify_audit_{timestamp}.xlsx"

    with csv_path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=DETAIL_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for row in report["details"]:
            safe_row = {key: _safe_spreadsheet_value(value) for key, value in row.items()}
            safe_row["findings"] = _safe_spreadsheet_value("; ".join(row.get("findings", [])))
            writer.writerow(safe_row)

    with employee_csv_path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=EMPLOYEE_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({key: _safe_spreadsheet_value(value) for key, value in row.items()} for row in report["employees"])

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Field", "Value"])
    metadata = [
        ("Generated at", report["generated_at"]),
        ("Timezone", report["timezone"]),
        ("Scan ID", report["scan_id"]),
        ("Scan captured at", report["scan_captured_at"]),
        ("Scan age hours", report["scan_age_hours"]),
        ("Scan stale", report["scan_stale"]),
        ("Pages scanned", report["pages_scanned"]),
        ("Snapshot products", report["snapshot_product_count"]),
    ]
    for row in metadata:
        summary_sheet.append([_safe_spreadsheet_value(value) for value in row])
    summary_sheet.append([])
    summary_sheet.append(["Audit metric", "Count"])
    for row in sorted(report["summary"].items()):
        summary_sheet.append(list(row))

    details_sheet = workbook.create_sheet("Details")
    details_sheet.append(DETAIL_FIELDS)
    for row in report["details"]:
        details_sheet.append([
            _safe_spreadsheet_value("; ".join(row.get(field, [])) if field == "findings" else row.get(field))
            for field in DETAIL_FIELDS
        ])

    employee_sheet = workbook.create_sheet("By Employee")
    employee_sheet.append(EMPLOYEE_FIELDS)
    for row in report["employees"]:
        employee_sheet.append([_safe_spreadsheet_value(row.get(field)) for field in EMPLOYEE_FIELDS])

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    workbook.save(xlsx_path)
    return {"csv": str(csv_path), "employees_csv": str(employee_csv_path), "xlsx": str(xlsx_path)}
