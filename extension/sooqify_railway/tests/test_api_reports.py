import csv
from datetime import datetime
from zoneinfo import ZoneInfo

import pytest

import app as server
from reporting import _employee_rows, build_daily_report, write_report_files


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("AUDIT_API_TOKEN", "test-audit-secret")
    monkeypatch.setattr(server, "PARTIAL_DIR", tmp_path / "partial")
    monkeypatch.setattr(server, "SCANS_DIR", tmp_path / "scans")
    monkeypatch.setattr(server, "SNAPSHOT_PATH", tmp_path / "latest.json")
    monkeypatch.setattr(server, "BASELINES_PATH", tmp_path / "baselines.json")
    monkeypatch.setattr(server, "REPORTS_DIR", tmp_path / "reports")
    return server.app.test_client()


def auth():
    return {"Authorization": "Bearer test-audit-secret"}


def test_employee_totals_exclude_unlinked_store_only_rows():
    rows = _employee_rows([
        {"local_id": 6017, "added_by": "Yusuf", "status": "match", "findings": []},
        {"local_id": None, "added_by": "غير محدد", "status": "unlinked_store", "findings": ["store_product_unlinked"]},
    ])

    assert rows == [{
        "employee": "Yusuf",
        "total": 1,
        "issues": 0,
        "unverified": 0,
        "unresolved": 0,
        "missing": 0,
        "name_mismatch": 0,
        "price_mismatch": 0,
        "brand_mismatch": 0,
        "store_id_mismatch": 0,
        "no_image": 0,
        "image_unknown": 0,
    }]


def test_scan_is_promoted_only_after_all_pages_arrive(client):
    scan_id = "scan-test-0001"
    started = client.post(
        "/api/audit/scans/start",
        headers=auth(),
        json={"scan_id": scan_id, "expected_pages": 2, "started_at": "2026-10-03T20:00:00+03:00"},
    )
    assert started.status_code == 200

    page_one = client.post(
        f"/api/audit/scans/{scan_id}/pages/1",
        headers=auth(),
        json={"products": [{
            "id": 6638,
            "name_ar": "ساعة باتيك فيليب",
            "brand_id": 5,
            "brand_name": "Patek Philippe",
            "price": 1320,
            "image_count": 1,
            "local_tags": ["6017"],
            "image_url": "https://store.example/private-path.png",
            "cookie": "must-not-be-stored",
        }]},
    )
    assert page_one.status_code == 200

    incomplete = client.post(
        f"/api/audit/scans/{scan_id}/complete",
        headers=auth(),
        json={"captured_at": "2026-10-03T20:10:00+03:00"},
    )
    assert incomplete.status_code == 409
    assert incomplete.get_json()["missing_pages"] == [2]
    assert not server.SNAPSHOT_PATH.exists()

    empty_page = client.post(
        f"/api/audit/scans/{scan_id}/pages/2",
        headers=auth(),
        json={"products": []},
    )
    assert empty_page.status_code == 200
    complete = client.post(
        f"/api/audit/scans/{scan_id}/complete",
        headers=auth(),
        json={"captured_at": "2026-10-03T20:10:00+03:00"},
    )
    assert complete.status_code == 200

    snapshot = server._load_json(server.SNAPSHOT_PATH)
    record = snapshot["products"][0]
    assert record["id"] == 6638
    assert record["local_tags"] == [6017]
    assert "image_url" not in record
    assert "cookie" not in record


def test_api_requires_token_and_health_does_not_reveal_secret(client):
    response = client.get("/api/audit/latest")
    assert response.status_code == 401
    health = client.get("/healthz")
    assert health.status_code == 200
    assert "test-audit-secret" not in health.get_data(as_text=True)


def test_pricing_baseline_preserves_original_cny_and_exact_fee_snapshot(client):
    response = client.post(
        "/api/pricing-baselines",
        headers=auth(),
        json={"products": [{
            "local_id": 6017,
            "name_en": "Watch",
            "name_ar": "ساعة",
            "product_type": "watches",
            "original_price_yuan": 500,
            "price_sar": 1320,
            "settings": {"ExchangeRate": 1.2, "WatchFlatFeeYuan": 600, "AddedFeeYuan": 250},
        }]},
    )
    assert response.status_code == 200
    baseline = server._load_price_baselines()["6017"]
    assert baseline["original_price_yuan"] == 500
    assert baseline["settings"] == {"ExchangeRate": 1.2, "WatchFlatFeeYuan": 600}


def test_report_generation_refuses_to_audit_against_an_unavailable_archive(client, tmp_path, monkeypatch):
    server._json_write_atomic(server.SNAPSHOT_PATH, {
        "complete": True,
        "scan_id": "scan-report-safe-01",
        "captured_at": "2026-10-03T20:10:00+03:00",
        "pages_scanned": 1,
        "product_count": 0,
        "products": [],
    })

    def unavailable_archive():
        raise OSError("offline")

    monkeypatch.setattr(server, "fetch_archive_snapshot", unavailable_archive)
    response = client.post("/api/reports/generate", headers=auth())

    assert response.status_code == 503
    assert "no audit report was generated" in response.get_json()["error"].lower()
    assert not list(server.REPORTS_DIR.glob("sooqify_audit_*"))


def test_reports_export_csv_and_xlsx_with_employee_sheet(tmp_path):
    snapshot = {
        "complete": True,
        "scan_id": "scan-test-0002",
        "captured_at": "2026-10-03T20:10:00+03:00",
        "pages_scanned": 1,
        "product_count": 1,
        "products": [{
            "id": 6638,
            "name_ar": "ساعة باتيك فيليب",
            "brand_id": 5,
            "brand_name": "Patek Philippe",
            "price": 1320,
            "image_count": None,
            "local_tags": ["6017"],
        }],
    }
    archive = {"sku": {
        "id": 6017,
        "name_en": "Patek Philippe Automatic Watch",
        "name_ar": "ساعة باتيك فيليب",
        "product_type": "watches",
        "brand_id": 5,
        "brand_name": "Patek Philippe",
        "price": 1320,
        "workflow_status": "submitted",
        "added_by": "Yusuf",
    }}
    report = build_daily_report(
        snapshot,
        archive,
        price_baselines={"6017": {
            "original_price_yuan": 500,
            "product_type": "watches",
            "settings": {"ExchangeRate": 1.2, "WatchFlatFeeYuan": 600},
        }},
        now=datetime(2026, 10, 3, 21, 0, tzinfo=ZoneInfo("Asia/Aden")),
    )
    paths = write_report_files(report, tmp_path)

    with open(paths["csv"], encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    assert rows[0]["local_id"] == "6017"
    assert rows[0]["expected_price_sar"] == "1320"
    assert report["employees"][0]["employee"] == "Yusuf"
    assert report["employees"][0]["image_unknown"] == 1
    assert paths["xlsx"].endswith(".xlsx")
    assert paths["employees_csv"].endswith("_employees.csv")
    with open(paths["employees_csv"], encoding="utf-8-sig", newline="") as file:
        employee_rows = list(csv.DictReader(file))
    assert employee_rows[0]["employee"] == "Yusuf"
    assert employee_rows[0]["total"] == "1"
    from openpyxl import load_workbook
    workbook = load_workbook(paths["xlsx"], read_only=True)
    assert "By Employee" in workbook.sheetnames
    assert "image_unknown" in next(workbook["By Employee"].iter_rows(values_only=True))


def test_sanitizer_preserves_generic_name_codes_and_unavailable_image_count(client):
    scan_id = "scan-generic-name-01"
    assert client.post(
        "/api/audit/scans/start", headers=auth(), json={"scan_id": scan_id, "expected_pages": 1}
    ).status_code == 200
    uploaded = client.post(
        f"/api/audit/scans/{scan_id}/pages/1",
        headers=auth(),
        json={"products": [{
            "id": 7010,
            "name": "Generic list label",
            "brand_name": "Value from an unavailable column",
            "brand_id": 9,
            "price": 123,
            "image_count": 3,
            "style_code": "STYLE-01",
            "search_code": "SEARCH-01",
            "fields_available": {
                "name": True,
                "price": False,
                "brand": False,
                "brand_id": False,
                "image_count": False,
                "style_code": True,
                "search_code": True,
                "local_id_tags": False,
            },
            "local_tags": ["6020"],
        }]},
    )
    assert uploaded.status_code == 200
    completed = client.post(f"/api/audit/scans/{scan_id}/complete", headers=auth(), json={})
    assert completed.status_code == 200

    record = server._load_json(server.SNAPSHOT_PATH)["products"][0]
    assert record["id"] == 7010
    assert record["name"] == "Generic list label"
    assert record["name_ar"] == ""
    assert record["name_en"] == ""
    assert record["image_count"] is None
    assert record["price"] is None
    assert record["brand_id"] is None
    assert record["brand_name"] == ""
    assert record["style_code"] == "STYLE-01"
    assert record["search_code"] == "SEARCH-01"
    assert record["fields_available"]["price"] is False
    assert record["fields_available"]["brand_id"] is False
    assert record["local_tags"] == []


def test_rejects_invalid_or_duplicate_store_ids(client):
    scan_id = "scan-test-0003"
    assert client.post("/api/audit/scans/start", headers=auth(), json={"scan_id": scan_id, "expected_pages": 1}).status_code == 200
    response = client.post(
        f"/api/audit/scans/{scan_id}/pages/1",
        headers=auth(),
        json={"products": [{"id": 0}, {"id": 1}]},
    )
    assert response.status_code == 400


def test_report_exports_escape_formula_like_employee_and_name_values(tmp_path):
    report = {
        "generated_at": "2026-10-03T21:00:00+03:00",
        "timezone": "Asia/Aden",
        "scan_id": "scan-safe-001",
        "scan_captured_at": "2026-10-03T20:00:00+03:00",
        "scan_age_hours": 1,
        "scan_stale": False,
        "pages_scanned": 1,
        "snapshot_product_count": 1,
        "summary": {"issues": 0},
        "details": [{"local_id": 1, "name_en": '=HYPERLINK("https://example.invalid")', "added_by": "@operator", "findings": []}],
        "employees": [{"employee": "=1+1", "total": 1, "issues": 0}],
    }
    paths = write_report_files(report, tmp_path)

    with open(paths["employees_csv"], encoding="utf-8-sig", newline="") as file:
        employee = next(csv.DictReader(file))
    assert employee["employee"] == "'=1+1"

    from openpyxl import load_workbook
    workbook = load_workbook(paths["xlsx"], read_only=True, data_only=False)
    details = next(workbook["Details"].iter_rows(min_row=2, values_only=True))
    employees = next(workbook["By Employee"].iter_rows(min_row=2, values_only=True))
    assert details[2].startswith("'=HYPERLINK")
    assert details[16] == "'@operator"
    assert employees[0] == "'=1+1"
